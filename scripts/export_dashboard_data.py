"""Lee el cuadro recalculado y genera dashboard/data.json (catálogo cap. 7 + elementos Revit)."""
import json, pathlib, sys
import openpyxl
root = pathlib.Path(__file__).resolve().parent.parent
wb = openpyxl.load_workbook(sys.argv[1] if len(sys.argv) > 1 else root / "CUADRO_CANTIDADES_UNIQUE76_RCI.xlsx", data_only=True)
nv, cat, grp, cap = wb["NIVEL"], [], "", False
for r in range(7, nv.max_row + 1):
    a, b, c = (nv.cell(r, i).value for i in (1, 2, 3))
    if b and str(b).startswith("COSTO DIRECTO"):
        break
    if a is not None and str(a).strip().isdigit():
        cap = str(a).strip() == "7"
        continue
    if not cap:
        continue
    if a and str(a).startswith("7."):
        cat.append([str(a), b.strip(), c or "Un.", grp])
    elif b:
        grp = b.strip()
rows = [[r[1], r[2] or "", r[5], round(float(r[7]), 2), r[8] or "", r[9] or "", r[10] or "", r[11] or "", r[13] or ""]
        for r in wb["BD_REVIT"].iter_rows(min_row=6, values_only=True) if r[0] is not None]
(root / "dashboard" / "data.json").write_text(json.dumps({"catalogo": cat, "rows": rows}, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
print(len(cat), "ítems,", len(rows), "elementos")
